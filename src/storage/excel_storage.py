"""
Excel storage using openpyxl for append-only ranking data.
"""

import logging
import hashlib
from pathlib import Path
from typing import List, Set, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.collectors.base_collector import RankingSnapshot
from .run_log import RunLog

logger = logging.getLogger(__name__)


class ExcelStorage:
    """
    Excel storage for ranking data with append-only sheets.

    Sheets:
    - rank_history_min: Ranking snapshots
    - run_log: Collection run logs

    Features:
    - Deduplication via deterministic row key
    - Append-only (never deletes existing data)
    - Automatic sheet creation with headers
    """

    # Column definitions for rank_history_min
    RANK_HISTORY_COLUMNS = [
        "date_kst",
        "market",
        "category_key",
        "category_url",
        "rank",
        "product_name_raw",
        "brand_raw",
        "product_url",
        "run_id",
        "parse_status",
        "notes"
    ]

    def __init__(
        self,
        output_file: str,
        rank_history_sheet: str = "rank_history_min",
        run_log_sheet: str = "run_log"
    ):
        """
        Initialize Excel storage.

        Args:
            output_file: Path to output Excel file
            rank_history_sheet: Name of ranking history sheet
            run_log_sheet: Name of run log sheet
        """
        self.output_file = Path(output_file)
        self.rank_history_sheet = rank_history_sheet
        self.run_log_sheet = run_log_sheet

        # Ensure directory exists
        self.output_file.parent.mkdir(parents=True, exist_ok=True)

        # Initialize workbook
        self._init_workbook()

        # Load existing row keys for deduplication
        self._existing_keys: Set[str] = set()
        self._load_existing_keys()

        logger.info(f"ExcelStorage initialized: {self.output_file}")

    def _init_workbook(self):
        """Initialize or load the workbook with required sheets."""
        if self.output_file.exists():
            self.workbook = load_workbook(self.output_file)
            logger.debug(f"Loaded existing workbook: {self.output_file}")
        else:
            self.workbook = Workbook()
            # Remove default sheet
            if "Sheet" in self.workbook.sheetnames:
                del self.workbook["Sheet"]
            logger.debug(f"Created new workbook: {self.output_file}")

        # Ensure rank_history_min sheet exists
        if self.rank_history_sheet not in self.workbook.sheetnames:
            ws = self.workbook.create_sheet(self.rank_history_sheet)
            ws.append(self.RANK_HISTORY_COLUMNS)
            logger.debug(f"Created sheet: {self.rank_history_sheet}")

        # Ensure run_log sheet exists
        if self.run_log_sheet not in self.workbook.sheetnames:
            ws = self.workbook.create_sheet(self.run_log_sheet)
            ws.append(RunLog.get_columns())
            logger.debug(f"Created sheet: {self.run_log_sheet}")

        self._save()

    def _save(self):
        """Save the workbook."""
        self.workbook.save(self.output_file)

    def _generate_row_key(self, snapshot: RankingSnapshot) -> str:
        """
        Generate a deterministic row key for deduplication.

        Key components: date_kst + market + category_key + rank + product_url

        Args:
            snapshot: Ranking snapshot

        Returns:
            Hash string for the row
        """
        key_parts = [
            snapshot.date_kst,
            snapshot.market,
            snapshot.category_key,
            str(snapshot.rank),
            snapshot.product_url
        ]
        key_string = "|".join(key_parts)
        return hashlib.md5(key_string.encode()).hexdigest()

    def _load_existing_keys(self):
        """Load existing row keys from the rank_history sheet."""
        if self.rank_history_sheet not in self.workbook.sheetnames:
            return

        ws = self.workbook[self.rank_history_sheet]

        # Find column indices
        header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]

        try:
            date_idx = header_row.index("date_kst")
            market_idx = header_row.index("market")
            category_idx = header_row.index("category_key")
            rank_idx = header_row.index("rank")
            url_idx = header_row.index("product_url")
        except ValueError:
            logger.warning("Could not find all key columns in existing sheet")
            return

        # Load keys from existing rows
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue

            key_parts = [
                str(row[date_idx] or ""),
                str(row[market_idx] or ""),
                str(row[category_idx] or ""),
                str(row[rank_idx] or ""),
                str(row[url_idx] or "")
            ]
            key_string = "|".join(key_parts)
            key_hash = hashlib.md5(key_string.encode()).hexdigest()
            self._existing_keys.add(key_hash)

        logger.debug(f"Loaded {len(self._existing_keys)} existing row keys")

    def _snapshot_to_row(self, snapshot: RankingSnapshot) -> list:
        """Convert a snapshot to a row list."""
        return [
            snapshot.date_kst,
            snapshot.market,
            snapshot.category_key,
            snapshot.category_url,
            snapshot.rank,
            snapshot.product_name_raw,
            snapshot.brand_raw or "",
            snapshot.product_url,
            snapshot.run_id,
            snapshot.parse_status,
            snapshot.notes
        ]

    def append_snapshots(self, snapshots: List[RankingSnapshot]) -> int:
        """
        Append ranking snapshots to the rank_history sheet.

        Deduplicates based on (date_kst, market, category_key, rank, product_url).

        Args:
            snapshots: List of RankingSnapshot objects

        Returns:
            Number of new rows added
        """
        if not snapshots:
            return 0

        ws = self.workbook[self.rank_history_sheet]
        added = 0

        for snapshot in snapshots:
            row_key = self._generate_row_key(snapshot)

            if row_key in self._existing_keys:
                logger.debug(f"Skipping duplicate: {snapshot.product_name_raw[:30]}...")
                continue

            row = self._snapshot_to_row(snapshot)
            ws.append(row)
            self._existing_keys.add(row_key)
            added += 1

        if added > 0:
            self._save()
            logger.info(f"Added {added} new ranking rows (skipped {len(snapshots) - added} duplicates)")

        return added

    def append_run_log(self, run_log: RunLog):
        """
        Append a run log entry to the run_log sheet.

        Args:
            run_log: RunLog object
        """
        ws = self.workbook[self.run_log_sheet]
        ws.append(run_log.to_row())
        self._save()
        logger.info(f"Added run log: {run_log.run_id}")

    def get_rank_history(
        self,
        market: Optional[str] = None,
        category_key: Optional[str] = None,
        days: Optional[int] = None
    ) -> List[dict]:
        """
        Retrieve ranking history with optional filters.

        Args:
            market: Filter by market (amazon_us, cosme_jp)
            category_key: Filter by category
            days: Limit to last N days (not implemented in MVP)

        Returns:
            List of ranking dictionaries
        """
        ws = self.workbook[self.rank_history_sheet]

        # Get header row
        header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]

        results = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue

            row_dict = dict(zip(header_row, row))

            # Apply filters
            if market and row_dict.get("market") != market:
                continue
            if category_key and row_dict.get("category_key") != category_key:
                continue

            results.append(row_dict)

        return results

    def get_unique_dates(self) -> List[str]:
        """Get list of unique dates in the rank history."""
        ws = self.workbook[self.rank_history_sheet]

        # Find date column
        header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        try:
            date_idx = header_row.index("date_kst")
        except ValueError:
            return []

        dates = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is not None and row[date_idx]:
                dates.add(str(row[date_idx]))

        return sorted(dates)

    def get_run_logs(self) -> List[dict]:
        """Retrieve all run logs."""
        ws = self.workbook[self.run_log_sheet]

        header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]

        results = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue

            row_dict = dict(zip(header_row, row))
            results.append(row_dict)

        return results

    def close(self):
        """Close the workbook."""
        if self.workbook:
            self.workbook.close()
            logger.debug("Workbook closed")
